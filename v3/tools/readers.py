# v3/tools/readers.py — document loading for ingest

from pathlib import Path

from langchain_core.documents import Document


def load_documents(filepath: str) -> list[Document]:
    path = Path(filepath)
    ext = path.suffix.lower()

    if ext in {".txt", ".md", ".csv", ".json", ".xml", ".html"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return [Document(page_content=text, metadata={"source": str(path.resolve())})]

    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = []
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(
                    Document(
                        page_content=text,
                        metadata={"source": str(path.resolve()), "page": index},
                    )
                )
        return pages or [Document(page_content="", metadata={"source": str(path.resolve())})]

    if ext == ".docx":
        from docx import Document as DocxDocument

        doc = DocxDocument(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n\n".join(paragraphs)
        return [Document(page_content=text, metadata={"source": str(path.resolve())})]

    if ext in {".xlsx", ".xls"}:
        return _load_spreadsheet(path)

    raise ValueError(f"Unsupported file type for ingest: {ext or 'unknown'}")


def _load_spreadsheet(path: Path) -> list[Document]:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    chunks: list[Document] = []

    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) if cell is not None else "" for cell in row]
            if any(cell.strip() for cell in cells):
                rows.append(" | ".join(cells))
        if rows:
            chunks.append(
                Document(
                    page_content=f"Sheet: {sheet.title}\n" + "\n".join(rows),
                    metadata={"source": str(path.resolve()), "sheet": sheet.title},
                )
            )

    workbook.close()
    return chunks or [Document(page_content="", metadata={"source": str(path.resolve())})]


def documents_to_text(documents: list[Document]) -> str:
    return "\n\n".join(doc.page_content for doc in documents if doc.page_content.strip())

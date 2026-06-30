# v3/tools/writers.py — DOCX output tool

from pathlib import Path


def write_docx(output_path: str, title: str, sections: dict[str, str]) -> str:
    from docx import Document
    from docx.shared import Pt

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    heading = doc.add_heading(title, level=0)
    heading.runs[0].font.size = Pt(20)

    for section_title, body in sections.items():
        doc.add_heading(section_title, level=1)
        for paragraph in body.split("\n"):
            stripped = paragraph.strip()
            if stripped:
                doc.add_paragraph(stripped)

    doc.save(str(path))
    return str(path.resolve())
